import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

def create_deck():
    prs = Presentation()
    # 16:9 widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Brand Colors
    AMETHYST = RGBColor(74, 21, 75)     # #4A154B (Primary Brand)
    TEAL = RGBColor(0, 128, 128)        # #008080 (Secondary Accent)
    GOLD = RGBColor(255, 193, 7)        # #FFC107 (Highlight / Hope)
    CRIMSON = RGBColor(211, 47, 47)     # #D32F2F (Alert)
    CHARCOAL = RGBColor(18, 18, 18)     # #121212 (Dark Base)
    CARD_BG = RGBColor(28, 28, 30)      # #1C1C1E (Card surface)
    CARD_BORDER = RGBColor(50, 50, 56)  # Subtle border
    WHITE = RGBColor(255, 255, 255)
    OFF_WHITE = RGBColor(240, 240, 245)
    MUTED = RGBColor(160, 160, 170)
    DARK_AMETHYST = RGBColor(40, 12, 42)

    def set_slide_background(slide, color=CHARCOAL):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, slide_num, title, subtitle):
        # Header banner shape
        banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.5), Inches(11.733), Inches(1.1))
        banner.fill.background()
        banner.line.fill.background()
        tf = banner.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        # Number pill / category
        p0 = tf.paragraphs[0]
        p0.text = f"SLIDE 0{slide_num}  •  DKC FELLOWSHIP 2026  •  THEME 04: GENDER-BASED ONLINE VIOLENCE"
        p0.font.name = "Arial"
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = GOLD

        # Title
        p1 = tf.add_paragraph()
        p1.text = title
        p1.font.name = "Arial"
        p1.font.size = Pt(24)
        p1.font.bold = True
        p1.font.color.rgb = WHITE

        # Subtitle
        if subtitle:
            p2 = tf.add_paragraph()
            p2.text = subtitle
            p2.font.name = "Arial"
            p2.font.size = Pt(12)
            p2.font.color.rgb = TEAL

    def create_card(slide, left, top, width, height, title="", border_color=CARD_BORDER, bg_color=CARD_BG):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)
        
        if title:
            tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), width - Inches(0.4), Inches(0.5))
            tf = tb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = title
            p.font.name = "Arial"
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = GOLD
        return card

    # ==========================================
    # SLIDE 1: Title & Hook
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, DARK_AMETHYST)

    # Accent decorative bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.0), Inches(0.15), Inches(5.5))
    bar.fill.solid()
    bar.fill.fore_color.rgb = GOLD
    bar.line.fill.background()

    # Title box
    tb1 = s1.shapes.add_textbox(Inches(1.2), Inches(0.8), Inches(7.5), Inches(5.8))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "DKC DIGITAL RESPECT & COHESION FELLOWSHIP 2026"
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = GOLD

    p = tf1.add_paragraph()
    p.text = "Project Protiti (প্রতীতি)"
    p.font.name = "Arial"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE

    p = tf1.add_paragraph()
    p.text = "“সন্দেহহীন সত্য, অকাট্য প্রমাণের অধিকার”\nFrom Doubt to Certitude: Structure for Digital Justice"
    p.font.name = "Arial"
    p.font.size = Pt(16)
    p.font.color.rgb = TEAL

    p = tf1.add_paragraph()
    p.text = "\nA Trauma-Informed Digital Evidence Vault, Legal Automator (Cyber Protection Act 2026), and Offline Emergency SOS Framework for Women in Bangladesh."
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.color.rgb = OFF_WHITE

    p = tf1.add_paragraph()
    p.text = "\nPartners: UNDP Bangladesh  •  European Union  •  SURGE\nTheme: Gender-based Online Violence  •  Lens: Gender Equity"
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.color.rgb = MUTED

    # Right Card: Team & Key Metric
    create_card(s1, Inches(9.0), Inches(0.8), Inches(3.6), Inches(5.9), title="FELLOWSHIP TEAM ROSTER", bg_color=CARD_BG, border_color=AMETHYST)
    tb_team = s1.shapes.add_textbox(Inches(9.2), Inches(1.4), Inches(3.2), Inches(5.1))
    tf_team = tb_team.text_frame
    tf_team.word_wrap = True

    team_content = [
        ("FARZANA REZA", "Team Leader, Operations, Marketing & Research (Acting UI/UX)", GOLD),
        ("ABRAR MASUD", "AI & Systems Engineer (Acting Full-Stack / Mobile)", TEAL),
        ("MITHILA KHAN", "Legal Advisor & Statutory Compliance (CPA 2026)", WHITE),
        ("RESERVE SLOT 4", "Dedicated UI/UX Designer (Interim: Farzana)", MUTED),
        ("RESERVE SLOT 5", "Dedicated Flutter Developer (Interim: Abrar)", MUTED),
    ]

    p_t0 = tf_team.paragraphs[0]
    p_t0.text = "CORE FELLOWS (Bootcamps & Leadership):"
    p_t0.font.size = Pt(10)
    p_t0.font.bold = True
    p_t0.font.color.rgb = WHITE

    for name, role, color in team_content:
        p_name = tf_team.add_paragraph()
        p_name.text = f"• {name}"
        p_name.font.size = Pt(11)
        p_name.font.bold = True
        p_name.font.color.rgb = color
        
        p_role = tf_team.add_paragraph()
        p_role.text = f"   {role}"
        p_role.font.size = Pt(9)
        p_role.font.color.rgb = OFF_WHITE

    p_stat = tf_team.add_paragraph()
    p_stat.text = "\nTHE CRITICAL REALITY:\n64% of Bangladeshi women face online harassment; 80%+ cases go unreported due to evidence destruction and legal intimidation."
    p_stat.font.size = Pt(9.5)
    p_stat.font.bold = True
    p_stat.font.color.rgb = CRIMSON

    # ==========================================
    # SLIDE 2: The Problem
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, 2, "The Reality: Online Violence Destroys Evidence & Silences Victims", 
               "Trauma, legal complexity, and social stigma prevent survivors from seeking justice.")

    col_width = Inches(2.75)
    top_pos = Inches(1.8)
    h_pos = Inches(5.0)

    cards_data = [
        ("1. Evidence Panic Deletion", 
         "CRIMSON",
         "Out of immediate panic or fear of family discovery, 80%+ of victims delete chat logs and block abusers, irreversibly erasing digital proof required by police and prosecutors.",
         "Impact: Without bitwise evidence, law enforcement cannot file an FIR under CrPC Sec 154."),
        ("2. Legal Intimidation", 
         "GOLD",
         "The Cyber Security Act has been repealed by the Cyber Protection Act 2026. General Diary (GD) drafting demands specific statutory mapping that ordinary survivors cannot navigate alone.",
         "Impact: Duty officers reject handwritten, unstructured victim complaints at local Thanas."),
        ("3. The Forensic Gap", 
         "TEAL",
         "Simple screenshots are easily dismissed in court. Defense lawyers routinely challenge them for lack of metadata, timestamps, URL origins, or Evidence Act Section 65B certificates.",
         "Impact: Forensic admissibility fails before the Divisional Cyber Tribunals."),
        ("4. Coercion & Stigma", 
         "AMETHYST",
         "Violent partners or household members regularly seize phones to check contents. Without decoy isolation, having an anti-harassment app invites immediate physical retribution.",
         "Impact: Victims remain trapped in silence under fear of physical assault.")
    ]

    for i, (title, color_name, desc, imp) in enumerate(cards_data):
        c_left = Inches(0.8) + i * Inches(2.95)
        border = CRIMSON if color_name == "CRIMSON" else GOLD if color_name == "GOLD" else TEAL if color_name == "TEAL" else AMETHYST
        create_card(s2, c_left, top_pos, col_width, h_pos, title=title, border_color=border)
        
        tb = s2.shapes.add_textbox(c_left + Inches(0.15), top_pos + Inches(0.7), col_width - Inches(0.3), h_pos - Inches(0.9))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = desc
        p.font.size = Pt(11)
        p.font.color.rgb = OFF_WHITE
        
        p2 = tf.add_paragraph()
        p2.text = f"\n{imp}"
        p2.font.size = Pt(10)
        p2.font.bold = True
        p2.font.color.rgb = GOLD if color_name != "GOLD" else TEAL

    # ==========================================
    # SLIDE 3: The Solution
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, 3, "The Solution: Protiti (প্রতীতি) — Structural Safety & Justice", 
               "An encrypted, trauma-informed digital legal assistant engineered for Bangladesh.")

    sol_cards = [
        ("🔒 On-Device Secure Vault",
         "• AES-256 / SQLCipher local encryption with zero cloud exposure of sensitive media.\n• Duress PIN (Decoy Vault): Entering 9999 loads benign notes/recipes to safely bypass physical coercion.\n• Stealth Calculator Disguise: Masked as functional arithmetic utility.",
         TEAL),
        ("📝 Automated e-GD Generator",
         "• Translates traumatic experiences into standard Bangladesh Police e-GD schema (gd.police.gov.bd).\n• Direct Section mapping to Cyber Protection Act 2026 & Pornography Control Act 2012.\n• Generates bilingual (Bengali/English) printable PDF with QR verification.",
         GOLD),
        ("🚨 Offline Panic SOS Frame",
         "• 3-Second Hold circular progress trigger with haptic feedback to prevent accidental triggers.\n• Zero-Data SMS Fallback: Broadcasts GPS coordinates to trusted contacts without cellular data.\n• One-touch speed dialing for National Emergency 999 and Police Cyber Support for Women (PCSW).",
         CRIMSON),
        ("🤝 Support Bridge Network",
         "• Direct institutional escalation to BLAST (2,500 panel lawyers) & BNWLA for pro-bono representation.\n• Integrated crisis counseling directory (Kaan Pete Roi, Moner Janala).\n• Confidential triage by vetted female cyber crime officers.",
         AMETHYST)
    ]

    for i, (title, content, color) in enumerate(sol_cards):
        col = i % 2
        row = i // 2
        c_left = Inches(0.8) + col * Inches(5.95)
        c_top = Inches(1.8) + row * Inches(2.55)
        
        create_card(s3, c_left, c_top, Inches(5.75), Inches(2.35), title=title, border_color=color)
        tb = s3.shapes.add_textbox(c_left + Inches(0.2), c_top + Inches(0.65), Inches(5.35), Inches(1.5))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = content
        p.font.size = Pt(11)
        p.font.color.rgb = OFF_WHITE

    # ==========================================
    # SLIDE 4: How It Works (Survivor Journey)
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, 4, "How It Works: 5-Stage Journey From Chaos to Certitude", 
               "Guided step-by-step pipeline transforming chaotic abuse into courtroom-admissible evidence.")

    steps = [
        ("Step 1: Capture", "Survivor captures screenshot, audio voice note, URL, or harassing chat before blocking abuser.", TEAL),
        ("Step 2: Authenticate", "Protiti extracts RFC 3161 timestamps, EXIF metadata, and generates bitwise SHA-256 forensic hash.", GOLD),
        ("Step 3: Triage", "Trauma-informed 4-step wizard guides survivor through emotional and factual details in plain Bangla.", AMETHYST),
        ("Step 4: Formulate", "Engine auto-maps facts to CPA 2026 sections, generating DMP-standard General Diary (e-GD payload).", TEAL),
        ("Step 5: Action", "Survivor submits directly to Thana / PCSW online, or dispatches securely to BLAST/BNWLA legal counsel.", CRIMSON),
    ]

    card_w = Inches(2.25)
    for i, (stitle, sdesc, scolor) in enumerate(steps):
        c_left = Inches(0.8) + i * Inches(2.4)
        create_card(s4, c_left, Inches(1.8), card_w, Inches(3.6), title=stitle, border_color=scolor)
        
        tb = s4.shapes.add_textbox(c_left + Inches(0.12), Inches(2.5), card_w - Inches(0.24), Inches(2.7))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = sdesc
        p.font.size = Pt(11)
        p.font.color.rgb = OFF_WHITE

    # Bottom comparison callout
    bot_card = create_card(s4, Inches(0.8), Inches(5.6), Inches(11.733), Inches(1.3), title="BEFORE vs AFTER PROTITI", border_color=GOLD)
    tb_b = s4.shapes.add_textbox(Inches(1.0), Inches(6.05), Inches(11.3), Inches(0.75))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True
    p = tf_b.paragraphs[0]
    p.text = "BEFORE: Panic deletion → 0 legal proof → Dismissed by police → Perpetrator escalates abuse.\nAFTER: Untouched encrypted vault → Bitwise SHA-256 verification → Evidence Act Sec 65B certified GD → Rapid police action."
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 5: Technology & Innovation
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, 5, "Technology & Security: Built for Real-World Survival", 
               "Engineered to withstand physical coercion, rural network blackouts, and courtroom scrutiny.")

    tech_pillars = [
        ("🔐 Dual-Vault & Duress PIN", 
         "Hardware-backed Keystore/Keychain encryption (SQLCipher AES-256). Entering Master PIN ('1234') unlocks forensic records; entering Duress PIN ('9999') loads an innocent personal diary.", 
         TEAL),
        ("📲 Launcher Stealth Disguise", 
         "Fully functioning 'Smart Calc' persona. Android PackageManager activity-alias toggling hides forensic markers from device inspectors. Face-down quick-exit gesture wipes active RAM.", 
         AMETHYST),
        ("📡 Zero-Data Offline SOS", 
         "Lightweight GPS extraction (<160 chars) transmitted over native cellular SMS without dependency on mobile data/Wi-Fi. Direct USSD & hotline fallbacks for rural divisional settings.", 
         CRIMSON),
        ("⚖️ Evidence Act Sec 65B Engine", 
         "SHA-256 cryptographic hashing (NIST FIPS 180-4) with automated Certificate of Electronic Authenticity adhering to the Evidence (Amendment) Act 2022 and CID forensic lab protocols.", 
         GOLD)
    ]

    for i, (title, text, color) in enumerate(tech_pillars):
        col = i % 2
        row = i // 2
        c_left = Inches(0.8) + col * Inches(5.95)
        c_top = Inches(1.8) + row * Inches(2.55)
        
        create_card(s5, c_left, c_top, Inches(5.75), Inches(2.35), title=title, border_color=color)
        tb = s5.shapes.add_textbox(c_left + Inches(0.2), c_top + Inches(0.65), Inches(5.35), Inches(1.5))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(11)
        p.font.color.rgb = OFF_WHITE

    # ==========================================
    # SLIDE 6: Target Community & Ecosystem Impact
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, 6, "Target Community & Measurable Ecosystem Impact", 
               "Bridging female university students, rural youth, police cyber cells, and pro-bono litigators.")

    # Left: Stakeholder cards
    create_card(s6, Inches(0.8), Inches(1.8), Inches(5.75), Inches(5.1), title="KEY BENEFICIARIES & STAKEHOLDERS", border_color=TEAL)
    tb_st = s6.shapes.add_textbox(Inches(1.0), Inches(2.4), Inches(5.35), Inches(4.3))
    tf_st = tb_st.text_frame
    tf_st.word_wrap = True

    stakeholders = [
        ("Primary Beneficiaries:", "Female university and college students (18–30) residing in student halls and dormitories facing digital harassment, non-consensual image distribution, and deepfakes."),
        ("Secondary Beneficiaries:", "Rural and peri-urban youth accessing digital safety toolkits through BRAC community networks and Union Digital Centres (UDCs)."),
        ("Institutional Stakeholders:", "Police Cyber Support for Women (PCSW), Thana cyber desks, and pro-bono panels at BLAST and BNWLA.")
    ]
    for idx, (head, body) in enumerate(stakeholders):
        p = tf_st.paragraphs[0] if idx == 0 else tf_st.add_paragraph()
        p.text = f"• {head} "
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = GOLD
        
        p2 = tf_st.add_paragraph()
        p2.text = f"   {body}\n"
        p2.font.size = Pt(10)
        p2.font.color.rgb = OFF_WHITE

    # Right: Quantitative Impact Targets
    create_card(s6, Inches(6.75), Inches(1.8), Inches(5.78), Inches(5.1), title="8-MONTH QUANTITATIVE IMPACT TARGETS", border_color=GOLD)
    tb_imp = s6.shapes.add_textbox(Inches(6.95), Inches(2.4), Inches(5.35), Inches(4.3))
    tf_imp = tb_imp.text_frame
    tf_imp.word_wrap = True

    impact_points = [
        ("80%+ Reduction in Evidence Destruction", "Survivors preserve bitwise intact records instead of deleting traumatic chats."),
        ("5,000 Discrete Toolkits Distributed", "Pocket-sized safety cards distributed across Dhaka Division campus dormitories and halls."),
        ("3x Increase in Admissible e-GDs Filed", "Eliminating bureaucratic rejection at local police stations via standardized schema."),
        ("100% Pro-Bono Legal Escalation", "Immediate referral to BLAST's 2,500 panel lawyers for indigent victims."),
        ("24/7 Offline Emergency Redundancy", "SMS-based coordinates ensure rural safety even in zero-data mobile blackouts.")
    ]
    for idx, (metric, subtext) in enumerate(impact_points):
        p = tf_imp.paragraphs[0] if idx == 0 else tf_imp.add_paragraph()
        p.text = f"✔ {metric}"
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = TEAL
        
        p2 = tf_imp.add_paragraph()
        p2.text = f"   {subtext}"
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = OFF_WHITE

    # ==========================================
    # SLIDE 7: Implementation Roadmap & Budget
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, 7, "8-Month Implementation Roadmap & BDT 50,000 Seed Grant", 
               "Strictly 0% spent on personal stipends; 100% dedicated to toolkits, workshops, and survivor safety.")

    # Left: Roadmap
    create_card(s7, Inches(0.8), Inches(1.8), Inches(5.75), Inches(5.1), title="8-MONTH PHASED TIMELINE", border_color=AMETHYST)
    tb_rd = s7.shapes.add_textbox(Inches(1.0), Inches(2.4), Inches(5.35), Inches(4.3))
    tf_rd = tb_rd.text_frame
    tf_rd.word_wrap = True

    phases = [
        ("Phase 1: Foundation & Audit (M1–M3)", "Oct – Dec 2026",
         "Finalize on-device SQLCipher vault, Duress PIN, e-GD schema. Conduct BLAST forensic hash validation. Attend Innovation Bootcamp."),
        ("Phase 2: Dhaka Division Pilot (M4–M6)", "Jan – Mar 2027",
         "Roll out 'Train the Trainer' workshops across Dhaka Division campus dorms (DU, JU, Eden Mohila College, BRACU, NSU). Distribute 5,000 survivor toolkits. Attend Progress Bootcamp."),
        ("Phase 3: Impact Showcase & National Roadmap (M7–M8)", "Apr – May 2027",
         "National showcase with Cyber Police Centre & UNDP presenting audited Dhaka Pilot results plus the National 8-Division Expansion Roadmap. Open-source the e-GD legal framework. Transition to alumni network.")
    ]
    for idx, (p_title, p_time, p_desc) in enumerate(phases):
        p = tf_rd.paragraphs[0] if idx == 0 else tf_rd.add_paragraph()
        p.text = f"• {p_title} [{p_time}]"
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = GOLD
        
        p2 = tf_rd.add_paragraph()
        p2.text = f"   {p_desc}\n"
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = OFF_WHITE

    # Right: BDT 50,000 Budget Breakdown (Dhaka Division Pilot — see budget_50k.md)
    create_card(s7, Inches(6.75), Inches(1.8), Inches(5.78), Inches(5.1), title="BDT 50,000 SEED GRANT — DHAKA DIVISION PILOT", border_color=GOLD)
    tb_bg = s7.shapes.add_textbox(Inches(6.95), Inches(2.4), Inches(5.35), Inches(4.3))
    tf_bg = tb_bg.text_frame
    tf_bg.word_wrap = True

    budget_items = [
        ("Survivor Toolkits & Handbooks (16%)", "BDT 8,000", "5,000 discrete folding survivor cards @ BDT 1.30 + 23 Ambassador/Partner Handbooks."),
        ("Campus Workshop Logistics (40%)", "BDT 20,000", "Refreshments, materials & signage for 10 campus activations across 5 Dhaka Division universities."),
        ("Ambassador Ground Logistics (24%)", "BDT 12,000", "Voucher-backed transit subsidies for 8 female student ambassadors, 8 months."),
        ("Hotline Triage Connectivity (12%)", "BDT 6,000", "Dedicated project SIMs, data packages & kiosk hotspots for emergency case triage."),
        ("Audit & National Roadmap Prep (8%)", "BDT 4,000", "Receipt archiving, compliance reporting, and the Phase 2 National Expansion Roadmap dossier.")
    ]
    for idx, (b_cat, b_amt, b_desc) in enumerate(budget_items):
        p = tf_bg.paragraphs[0] if idx == 0 else tf_bg.add_paragraph()
        p.text = f"• {b_cat} — {b_amt}"
        p.font.bold = True
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEAL
        
        p2 = tf_bg.add_paragraph()
        p2.text = f"   {b_desc}"
        p2.font.size = Pt(9)
        p2.font.color.rgb = OFF_WHITE

    p_rule = tf_bg.add_paragraph()
    p_rule.text = "\n*UNDP / DKC COMPLIANCE: Disbursed in 2 tranches (BDT 25k x 2). Strictly BDT 0 spent on personal stipends or honoraria."
    p_rule.font.size = Pt(8.5)
    p_rule.font.bold = True
    p_rule.font.color.rgb = GOLD

    # ==========================================
    # SLIDE 8: Sustainability & Call to Action
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8, DARK_AMETHYST)

    # Accent decorative bar
    bar8 = s8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.0), Inches(0.15), Inches(5.5))
    bar8.fill.solid()
    bar8.fill.fore_color.rgb = GOLD
    bar8.line.fill.background()

    tb8 = s8.shapes.add_textbox(Inches(1.2), Inches(0.8), Inches(6.5), Inches(5.8))
    tf8 = tb8.text_frame
    tf8.word_wrap = True

    p = tf8.paragraphs[0]
    p.text = "SUSTAINABILITY & INSTITUTIONAL HANDOVER"
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = GOLD

    p = tf8.add_paragraph()
    p.text = "A Lasting Frame for Digital Justice"
    p.font.name = "Arial"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = WHITE

    sustain_points = [
        ("Institutional Handover", "Engineered for long-term stewardship by Police Cyber Support for Women (PCSW) and the ICT Division's Cyber Help initiatives."),
        ("Open-Source Legal Schema", "The e-GD mapping engine is released as open-source civic tech, allowing civic organizations to update legal sections as laws evolve."),
        ("Alumni Ambassador Continuity", "Trained student ambassadors join the permanent DKC Alumni Network to train incoming university cohorts.")
    ]
    for title, desc in sustain_points:
        p = tf8.add_paragraph()
        p.text = f"\n✔ {title}"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = TEAL
        
        p = tf8.add_paragraph()
        p.text = desc
        p.font.size = Pt(11)
        p.font.color.rgb = OFF_WHITE

    # Right Card: The Call to Action
    create_card(s8, Inches(8.0), Inches(1.2), Inches(4.5), Inches(5.0), title="THE CALL TO ACTION", border_color=GOLD, bg_color=CARD_BG)
    tb_cta = s8.shapes.add_textbox(Inches(8.3), Inches(2.0), Inches(3.9), Inches(3.8))
    tf_cta = tb_cta.text_frame
    tf_cta.word_wrap = True

    p_quote = tf_cta.paragraphs[0]
    p_quote.text = "“Every woman in Bangladesh deserves a frame of safety. Together, we build the structure for undeniable certitude.”"
    p_quote.font.size = Pt(16)
    p_quote.font.bold = True
    p_quote.font.italic = True
    p_quote.font.color.rgb = GOLD

    p_close = tf_cta.add_paragraph()
    p_close.text = "\n\nProject Protiti (প্রতীতি)\nDigital Khichuri Challenge 2026\nUNDP Bangladesh • European Union • SURGE"
    p_close.font.size = Pt(12)
    p_close.font.color.rgb = WHITE

    p_contact = tf_cta.add_paragraph()
    p_contact.text = "\nTeam Lead: Farzana Reza\nReady for Stage 1 Evaluation & Innovation Bootcamp."
    p_contact.font.size = Pt(11)
    p_contact.font.bold = True
    p_contact.font.color.rgb = TEAL

    # Save output
    output_path = "/Users/blackbird/Everything/dev/DKC/pitch_materials/Protiti_Pitch_Deck.pptx"
    prs.save(output_path)
    print(f"Presentation successfully saved to: {output_path}")

if __name__ == "__main__":
    create_deck()
